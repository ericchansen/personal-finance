"""Parse Vanguard ``customActivityReport`` Excel workbooks.

Unlike Vanguard's combined CSV export, these workbooks contain transaction
history for one account.  The account number is only present in the download
filename, so it is deliberately inferred there rather than guessed from data.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


HEADERS = (
    "Settlement date",
    "Trade date",
    "Symbol",
    "Name",
    "Type",
    "Account type",
    "Quantity",
    "Price",
    "Commission & fees**",
    "Amount",
)
_NORMALIZED_HEADERS = tuple(header.casefold() for header in HEADERS)
_ACCOUNT_FILENAME = re.compile(r"^customActivityReport\s+(.+)$", re.IGNORECASE)
_DESCRIPTIVE_ACCOUNT_FILENAME = re.compile(r"\(([^()]+)\)")
_DATE_FORMATS = ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d")


class VanguardActivityError(ValueError):
    """The workbook looks relevant but contains data that cannot be trusted."""


@dataclass(frozen=True)
class VanguardAccountIdentity:
    account_number: str
    account_types: tuple[str, ...]


@dataclass(frozen=True)
class VanguardActivity:
    transaction_date: date
    settlement_date: date | None
    holding: str
    symbol: str | None
    transaction_type: str
    shares: Decimal | None
    share_price: Decimal | None
    cash_amount: Decimal
    fees: Decimal


@dataclass(frozen=True)
class VanguardActivityReport:
    source: str
    account: VanguardAccountIdentity
    transactions: tuple[VanguardActivity, ...]
    duplicate_count: int = 0


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized_row(values: Iterable[Any]) -> tuple[str, ...]:
    return tuple(_text(value).casefold() for value in values)


def _header_index(worksheet: Any) -> tuple[int, dict[str, int]] | None:
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 100), values_only=True),
        start=1,
    ):
        normalized = _normalized_row(row)
        positions = {value: index for index, value in enumerate(normalized) if value}
        if all(header in positions for header in _NORMALIZED_HEADERS):
            return row_number, {
                original: positions[normalized_header]
                for original, normalized_header in zip(HEADERS, _NORMALIZED_HEADERS)
            }
    return None


def sniff_workbook(path: str | Path) -> bool:
    """Identify the export from workbook headers, not its filename."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, ValueError):
        return False
    try:
        return any(_header_index(sheet) is not None for sheet in workbook.worksheets)
    finally:
        workbook.close()


def _account_number(path: Path, account_number: str | None) -> str:
    if account_number is not None:
        result = account_number.strip()
    else:
        match = _ACCOUNT_FILENAME.fullmatch(path.stem.strip())
        if match:
            result = match.group(1).strip()
        else:
            descriptive = _DESCRIPTIVE_ACCOUNT_FILENAME.findall(path.stem)
            result = descriptive[-1].strip() if descriptive else ""
    if not result:
        raise VanguardActivityError(
            "account number is absent; use a 'customActivityReport <account>.xlsx' "
            "filename or pass account_number"
        )
    return result


def _date(value: Any, *, epoch: datetime, field: str, location: str) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch)
        except (TypeError, ValueError, OverflowError) as exc:
            raise VanguardActivityError(f"{location}: invalid {field}") from exc
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
        raise VanguardActivityError(f"{location}: invalid {field}")
    text = _text(value)
    if text.startswith("="):
        raise VanguardActivityError(f"{location}: formulas are not supported in {field}")
    for date_format in _DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            pass
    raise VanguardActivityError(f"{location}: invalid {field}")


def _decimal(
    value: Any,
    *,
    field: str,
    location: str,
    required: bool = False,
    free_is_zero: bool = False,
) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise VanguardActivityError(f"{location}: missing {field}")
        return None
    if isinstance(value, bool):
        raise VanguardActivityError(f"{location}: invalid {field}")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))

    text = _text(value)
    if text.startswith("="):
        raise VanguardActivityError(f"{location}: formulas are not supported in {field}")
    if free_is_zero and text.casefold() == "free":
        return Decimal("0")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = text.replace("$", "").replace(",", "").replace("+", "").strip()
    try:
        parsed = Decimal(text)
    except InvalidOperation as exc:
        raise VanguardActivityError(f"{location}: invalid {field}") from exc
    if not parsed.is_finite():
        raise VanguardActivityError(f"{location}: invalid {field}")
    return -parsed if negative else parsed


def _required_text(value: Any, *, field: str, location: str) -> str:
    text = _text(value)
    if not text:
        raise VanguardActivityError(f"{location}: missing {field}")
    if text.startswith("="):
        raise VanguardActivityError(f"{location}: formulas are not supported in {field}")
    return text


def _transaction(
    row: tuple[Any, ...],
    columns: dict[str, int],
    *,
    epoch: datetime,
    location: str,
) -> tuple[VanguardActivity, str]:
    def cell(header: str) -> Any:
        index = columns[header]
        return row[index] if index < len(row) else None

    transaction_date = _date(
        cell("Trade date"), epoch=epoch, field="trade date", location=location
    )
    if transaction_date is None:
        raise VanguardActivityError(f"{location}: missing trade date")

    fees = _decimal(
        cell("Commission & fees**"),
        field="fees",
        location=location,
        free_is_zero=True,
    )
    return (
        VanguardActivity(
            transaction_date=transaction_date,
            settlement_date=_date(
                cell("Settlement date"),
                epoch=epoch,
                field="settlement date",
                location=location,
            ),
            holding=_required_text(cell("Name"), field="holding", location=location),
            symbol=_text(cell("Symbol")) or None,
            transaction_type=_required_text(
                cell("Type"), field="transaction type", location=location
            ),
            shares=_decimal(cell("Quantity"), field="shares", location=location),
            share_price=_decimal(cell("Price"), field="share price", location=location),
            cash_amount=_decimal(
                cell("Amount"), field="cash amount", location=location, required=True
            ),
            fees=fees if fees is not None else Decimal("0"),
        ),
        _text(cell("Account type")),
    )


def parse_file(
    path: str | Path, *, account_number: str | None = None
) -> VanguardActivityReport:
    """Parse every matching sheet and remove exact duplicate transactions."""
    workbook_path = Path(path)
    number = _account_number(workbook_path, account_number)
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except (OSError, ValueError) as exc:
        raise VanguardActivityError(f"cannot read workbook: {workbook_path.name}") from exc

    transactions: list[VanguardActivity] = []
    account_types: set[str] = set()
    duplicates = 0
    seen: set[VanguardActivity] = set()
    matched_sheet = False
    try:
        for sheet in workbook.worksheets:
            header = _header_index(sheet)
            if header is None:
                continue
            matched_sheet = True
            header_row, columns = header
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                populated = [value for value in row if _text(value)]
                if not populated:
                    continue
                # Vanguard appends merged legal-notice rows below the table.
                if len(populated) == 1:
                    continue
                location = f"{sheet.title}!row {row_number}"
                transaction, account_type = _transaction(
                    row, columns, epoch=workbook.epoch, location=location
                )
                if account_type:
                    account_types.add(account_type)
                if transaction in seen:
                    duplicates += 1
                    continue
                seen.add(transaction)
                transactions.append(transaction)
    finally:
        workbook.close()

    if not matched_sheet:
        raise VanguardActivityError("workbook does not contain Vanguard activity headers")
    return VanguardActivityReport(
        source=workbook_path.name,
        account=VanguardAccountIdentity(number, tuple(sorted(account_types))),
        transactions=tuple(transactions),
        duplicate_count=duplicates,
    )


def parse_files(paths: Iterable[str | Path]) -> tuple[VanguardActivityReport, ...]:
    """Parse reports, removing overlap between files for the same account."""
    reports: list[VanguardActivityReport] = []
    seen_by_account: dict[str, set[VanguardActivity]] = {}
    for path in paths:
        report = parse_file(path)
        seen = seen_by_account.setdefault(report.account.account_number, set())
        unique: list[VanguardActivity] = []
        duplicates = report.duplicate_count
        for transaction in report.transactions:
            if transaction in seen:
                duplicates += 1
                continue
            seen.add(transaction)
            unique.append(transaction)
        reports.append(
            VanguardActivityReport(
                source=report.source,
                account=report.account,
                transactions=tuple(unique),
                duplicate_count=duplicates,
            )
        )
    return tuple(reports)
