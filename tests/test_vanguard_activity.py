from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from importers.extracts import vanguard_activity


HEADERS = [
    "Settlement date",
    "Trade date",
    "Symbol",
    "Name",
    "Type",
    "Account type",
    "Quantity",
    "Price",
    "Commission & fees**",
    "Amount",
]


def _workbook(tmp_path, name="customActivityReport 00009999.xlsx", rows=(), setup=None):
    path = tmp_path / name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transaction History"
    sheet.append(["Fictional custom report"])
    sheet.append([])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    if setup:
        setup(workbook)
    workbook.save(path)
    return path


def _row(**overrides):
    values = {
        "settlement": "8/21/2026",
        "trade": "08/20/2026",
        "symbol": "FICT",
        "name": "Fictional Index Fund",
        "kind": "Purchase",
        "account_type": "Fictional Brokerage",
        "quantity": "1,234.5678",
        "price": "$12.3400",
        "fees": "Free",
        "amount": "($15,234.56)",
    }
    values.update(overrides)
    return list(values.values())


def test_parses_identity_dates_and_decimal_formats(tmp_path):
    path = _workbook(tmp_path, rows=[_row()])

    report = vanguard_activity.parse_file(path)

    assert report.account == vanguard_activity.VanguardAccountIdentity(
        "00009999", ("Fictional Brokerage",)
    )
    assert report.transactions == (
        vanguard_activity.VanguardActivity(
            transaction_date=date(2026, 8, 20),
            settlement_date=date(2026, 8, 21),
            holding="Fictional Index Fund",
            symbol="FICT",
            transaction_type="Purchase",
            shares=Decimal("1234.5678"),
            share_price=Decimal("12.3400"),
            cash_amount=Decimal("-15234.56"),
            fees=Decimal("0"),
        ),
    )


def test_blank_rows_and_cash_without_symbol_are_supported(tmp_path):
    path = _workbook(
        tmp_path,
        rows=[
            [None] * 10,
            _row(
                settlement=None,
                symbol=None,
                name="Fictional Cash Reserve",
                kind="Cash distribution",
                quantity=None,
                price=None,
                fees=None,
                amount="$42.00",
            ),
        ],
    )

    transaction = vanguard_activity.parse_file(path).transactions[0]

    assert transaction.symbol is None
    assert transaction.settlement_date is None
    assert transaction.shares is None
    assert transaction.share_price is None
    assert transaction.fees == Decimal("0")


def test_native_excel_date_cells_are_normalized(tmp_path):
    path = _workbook(
        tmp_path,
        rows=[_row(settlement=date(2026, 7, 2), trade=datetime(2026, 7, 1, 14, 30))],
    )

    transaction = vanguard_activity.parse_file(path).transactions[0]

    assert transaction.transaction_date == date(2026, 7, 1)
    assert transaction.settlement_date == date(2026, 7, 2)


def test_unknown_transaction_type_is_preserved_verbatim(tmp_path):
    path = _workbook(tmp_path, rows=[_row(kind="Fictional special adjustment")])

    assert (
        vanguard_activity.parse_file(path).transactions[0].transaction_type
        == "Fictional special adjustment"
    )


def test_matching_sheet_is_found_among_multiple_sheets(tmp_path):
    def setup(workbook):
        workbook["Transaction History"].title = "Read Me"
        activity = workbook.create_sheet("Activity Data")
        activity.append(["Introduction"])
        activity.append(HEADERS)
        activity.append(_row())
        calculations = workbook.create_sheet("Calculations")
        calculations["A1"] = "=1+1"

    path = _workbook(tmp_path, setup=setup)

    assert vanguard_activity.sniff_workbook(path)
    assert len(vanguard_activity.parse_file(path).transactions) == 1


def test_sniffing_uses_headers_not_filename(tmp_path):
    valid = _workbook(tmp_path, name="renamed.xlsx", rows=[_row()])
    invalid = tmp_path / "customActivityReport 12345678.xlsx"
    workbook = Workbook()
    workbook.active.append(["Unrelated", "Workbook"])
    workbook.save(invalid)

    assert vanguard_activity.sniff_workbook(valid)
    assert not vanguard_activity.sniff_workbook(invalid)
    assert (
        vanguard_activity.parse_file(valid, account_number="fictional-account")
        .account.account_number
        == "fictional-account"
    )


def test_descriptive_archived_filename_preserves_account_identity(tmp_path):
    path = _workbook(
        tmp_path,
        name=(
            "Transaction History - Vanguard Fictional IRA (00009999) "
            "- 2020-01-01 to 2026-01-01.xlsx"
        ),
        rows=[_row()],
    )
    assert vanguard_activity.parse_file(path).account.account_number == "00009999"


def test_duplicate_rows_and_overlapping_files_are_detected(tmp_path):
    first_dir = tmp_path / "first"
    second_dir = tmp_path / "second"
    first_dir.mkdir()
    second_dir.mkdir()
    first = _workbook(
        first_dir, name="customActivityReport 00001111.xlsx", rows=[_row(), _row()]
    )
    second = _workbook(
        second_dir, name="customActivityReport 00001111.xlsx", rows=[_row()]
    )

    within = vanguard_activity.parse_file(first)
    combined = vanguard_activity.parse_files([first, second])

    assert len(within.transactions) == 1
    assert within.duplicate_count == 1
    assert len(combined[1].transactions) == 0
    assert combined[1].duplicate_count == 1


@pytest.mark.parametrize("value", [None, "", "not money", "NaN", "=SUM(A1:A2)"])
def test_missing_malformed_and_formula_amounts_are_rejected(tmp_path, value):
    path = _workbook(tmp_path, rows=[_row(amount=value)])

    with pytest.raises(vanguard_activity.VanguardActivityError, match="cash amount"):
        vanguard_activity.parse_file(path)


def test_formula_in_optional_numeric_field_is_rejected(tmp_path):
    path = _workbook(tmp_path, rows=[_row(quantity="=1+1")])

    with pytest.raises(vanguard_activity.VanguardActivityError, match="formulas"):
        vanguard_activity.parse_file(path)


def test_malformed_transaction_date_is_rejected(tmp_path):
    path = _workbook(tmp_path, rows=[_row(trade="not a date")])

    with pytest.raises(vanguard_activity.VanguardActivityError, match="trade date"):
        vanguard_activity.parse_file(path)
